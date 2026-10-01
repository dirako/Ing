"""Expedientes públicos SECOP; extracción trazable y borradores sin completar datos ausentes.

No ejecuta JavaScript, no inicia sesión, no resuelve CAPTCHA ni presenta ofertas.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import io
import json
import os
import re
import uuid
import zipfile
from datetime import datetime, timezone
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from html.parser import HTMLParser
from pathlib import Path
from urllib.error import HTTPError, URLError
from urllib.parse import urljoin, urlsplit
from urllib.request import Request, HTTPRedirectHandler, build_opener
from xml.etree import ElementTree as ET

from secop_query import _fold

ROOT = Path(__file__).resolve().parents[1]
PROFILE = json.loads((ROOT / "config/profile.json").read_text(encoding="utf-8"))
LIMITS = PROFILE["documentos"]
MAX_BYTES = LIMITS["max_bytes"]
MAX_DOCUMENTS = LIMITS["max_documentos"]
HOSTS = tuple(LIMITS["dominios_publicos"])
PENDING = "PENDIENTE DE INFORMACIÓN"


def now():
    return datetime.now(timezone.utc).isoformat()


def storage():
    return Path(os.environ.get("SECOP_OUTPUT_DIR") or str(ROOT / "expedientes")).resolve()


def folder(expediente: str):
    if not re.fullmatch(r"[a-f0-9]{32}", expediente):
        raise ValueError("Identificador de expediente inválido")
    path = (storage() / expediente).resolve()
    if not path.is_relative_to(storage()):
        raise ValueError("Ruta fuera del directorio de expedientes")
    return path


def save(path, value):
    temp = path.with_suffix(path.suffix + ".tmp")
    temp.write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding="utf-8")
    temp.replace(path)


def validate_url(url):
    parts = urlsplit(url)
    host = (parts.hostname or "").lower()
    if (parts.scheme != "https" or parts.username or parts.password or parts.port not in (None, 443)
            or not any(host == h or host.endswith("." + h) for h in HOSTS)):
        raise ValueError("URL no permitida: se requiere HTTPS de un dominio público SECOP configurado")
    return url


class PublicRedirect(HTTPRedirectHandler):
    def redirect_request(self, req, fp, code, msg, headers, newurl):
        validate_url(newurl)
        return super().redirect_request(req, fp, code, msg, headers, newurl)


def fetch(url):
    validate_url(url)
    request = Request(url, headers={"User-Agent": "SECOP-Intelligence/0.2", "Accept-Encoding": "identity"})
    with build_opener(PublicRedirect()).open(request, timeout=LIMITS["timeout_segundos"]) as response:
        validate_url(response.url)
        data = response.read(MAX_BYTES + 1)
        if len(data) > MAX_BYTES:
            raise ValueError("Documento excede el límite de bytes configurado")
        return data, response.headers.get_content_type(), response.url


class Links(HTMLParser):
    def __init__(self):
        super().__init__()
        self.links = []

    def handle_starttag(self, tag, attrs):
        attrs = dict(attrs)
        if tag == "a" and attrs.get("href"):
            self.links.append(attrs["href"])


def html_response(data, content_type):
    head = data[:2048].lstrip().lower()
    return "html" in content_type or head.startswith((b"<!doctype html", b"<html"))


def failure(exc):
    if isinstance(exc, HTTPError):
        return {"estado": "RESTRINGIDO" if exc.code in (401, 403, 429) else "ERROR_HTTP",
                "http_status": exc.code, "motivo": str(exc)}
    return {"estado": "NO_DESCARGADO", "motivo": str(exc)}


def discover(url):
    """Descubre enlaces estáticos; no afirma que el inventario del portal esté completo."""
    try:
        data, mime, final = fetch(url)
        if not html_response(data, mime):
            return {"estado": "ENLACE_DIRECTO", "url": url, "documentos": [url]}
        parser = Links()
        text = data.decode("utf-8", errors="replace")
        parser.feed(text)
        links = []
        for href in parser.links:
            target = urljoin(final, href).split("#")[0]
            if re.search(r"\.(pdf|docx?|xlsx?|csv|txt|zip)(?:[?]|$)|download|descarg|GetDocument", target, re.I):
                if target not in links:
                    links.append(target)
        blocked = re.search(r"captcha|access denied|acceso denegado|type=[\"']password", text, re.I) or "/login" in final.lower()
        return {"estado": "RESTRINGIDO" if blocked and not links else "INVENTARIO_PARCIAL",
                "url": url, "url_final": final, "documentos": links,
                "aviso": "Solo enlaces públicos estáticos. Puede requerir navegador, sesión o descarga manual; no se garantiza inventario completo."}
    except (OSError, ValueError, URLError) as exc:
        return {"url": url, "documentos": [], **failure(exc)}


def file_kind(data, mime):
    if html_response(data, mime):
        raise ValueError("El portal devolvió HTML (posible sesión, CAPTCHA o página intermedia), no el documento")
    if data.startswith(b"%PDF-"):
        return ".pdf"
    if data.startswith(b"PK"):
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            names = archive.namelist()
            if "word/document.xml" in names:
                return ".docx"
            if "xl/workbook.xml" in names:
                return ".xlsx"
        return ".zip"
    if mime.startswith("text/") or mime in ("application/csv", "application/json"):
        data.decode("utf-8-sig")
        return ".txt"
    return ".bin"


def download(url, directory, index):
    record = {"id": f"D{index:03}", "url": url, "consultado_en": now()}
    try:
        data, mime, final = fetch(url)
        suffix = file_kind(data, mime)
        name = record["id"] + suffix
        (directory / name).write_bytes(data)
        record.update(estado="DESCARGADO", archivo=name, url_final=final, mime=mime,
                      bytes=len(data), sha256=hashlib.sha256(data).hexdigest())
    except (OSError, ValueError, URLError, zipfile.BadZipFile) as exc:
        record.update(failure(exc))
    return record


def obtain(proceso: dict, urls: list[str] | None = None):
    if proceso.get("fuente") not in ("SECOP I", "SECOP II"):
        raise ValueError("Indique fuente SECOP I o SECOP II")
    if not (proceso.get("id_proceso") or proceso.get("id_contrato") or proceso.get("url")):
        raise ValueError("Se requiere identificador o URL del proceso/contrato seleccionado")
    expediente = uuid.uuid4().hex
    directory = folder(expediente)
    directory.mkdir(parents=True)
    discovery = discover(proceso["url"]) if proceso.get("url") else {
        "estado": "SIN_URL", "documentos": [], "aviso": "No se proporcionó enlace público del proceso"}
    links = list(dict.fromkeys((urls or []) + discovery["documentos"]))
    documents = [download(url, directory, i + 1) for i, url in enumerate(links[:MAX_DOCUMENTS])]
    for i, url in enumerate(links[MAX_DOCUMENTS:], start=MAX_DOCUMENTS + 1):
        documents.append({"id": f"D{i:03}", "url": url, "estado": "OMITIDO_LIMITE",
                          "motivo": "Máximo de documentos por expediente alcanzado"})
    manifest = {"expediente": expediente, "creado_en": now(), "proceso": proceso,
                "descubrimiento": discovery, "documentos": documents,
                "inventario_completo_verificado": False}
    save(directory / "manifest.json", manifest)
    return manifest


def zip_read(archive, name):
    if archive.getinfo(name).file_size > MAX_BYTES:
        raise ValueError("Contenido descomprimido excede el límite")
    return archive.read(name)


def extract(path):
    """Devuelve unidades citables; las páginas sin texto se conservan como pendientes OCR."""
    if path.suffix == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(path)
        if reader.is_encrypted:
            raise ValueError("PDF cifrado; requiere copia pública legible")
        return [{"ubicacion": f"página {i + 1}", "texto": page.extract_text() or ""}
                for i, page in enumerate(reader.pages)]
    if path.suffix == ".docx":
        with zipfile.ZipFile(path) as archive:
            root = ET.fromstring(zip_read(archive, "word/document.xml"))
        ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
        return [{"ubicacion": f"párrafo {i + 1} (incluye tablas)",
                 "texto": "".join(p.itertext()) if not p.findall(".//w:t", ns) else
                 "".join(t.text or "" for t in p.findall(".//w:t", ns))}
                for i, p in enumerate(root.findall(".//w:p", ns))]
    if path.suffix == ".xlsx":
        from openpyxl import load_workbook
        with zipfile.ZipFile(path) as archive:
            if sum(info.file_size for info in archive.infolist()) > MAX_BYTES:
                raise ValueError("XLSX descomprimido excede el límite")
        book = load_workbook(path, read_only=True, data_only=False, keep_links=False)
        try:
            return [{"ubicacion": f"hoja {sheet.title}, fila {row[0].row}",
                     "texto": " | ".join(f"{c.coordinate}: {c.value}" for c in row if c.value is not None)}
                    for sheet in book for row in sheet.iter_rows() if any(c.value is not None for c in row)]
        finally:
            book.close()
    if path.suffix == ".txt":
        return [{"ubicacion": f"línea {i + 1}", "texto": text}
                for i, text in enumerate(path.read_text(encoding="utf-8-sig").splitlines())]
    raise ValueError("Formato no extraído automáticamente; convertir XLS/DOC/ZIP o aportar texto verificado")


CATEGORIES = {
    "juridico": r"habilit|juridic|registro unico|\brup\b|certific|representante|inhabilidad",
    "financiero": r"financier|liquidez|endeudamiento|capital de trabajo|patrimonio|rentabilidad",
    "tecnico": r"tecnic|experiencia|equipo|especificaci|entregable|debera|requisito|obligaci",
    "economico": r"precio|presupuesto|cantidad|impuesto|\biva\b|moneda|valor unitario",
    "comercial": r"pago|garantia|vigencia de la oferta|validez|entrega|plazo",
    "evaluacion": r"puntaje|puntuaci|evaluaci|criterio|ponderaci|desempate",
    "cronograma": r"cronograma|fecha|cierre|adenda|audiencia",
}


def analyze(expediente):
    directory = folder(expediente)
    manifest = json.loads((directory / "manifest.json").read_text(encoding="utf-8"))
    requirements, texts, gaps = [], [], []
    for doc in manifest["documentos"]:
        if doc["estado"] != "DESCARGADO":
            gaps.append({"documento": doc["id"], "motivo": doc.get("motivo", doc["estado"])})
            continue
        path = (directory / doc["archivo"]).resolve()
        try:
            if not path.is_relative_to(directory) or hashlib.sha256(path.read_bytes()).hexdigest() != doc["sha256"]:
                raise ValueError("Archivo fuera del expediente o huella SHA256 distinta del manifiesto")
            units = extract(path)
            texts.append({"documento": doc["id"], "unidades": units})
            if not units or not any(u["texto"].strip() for u in units):
                gaps.append({"documento": doc["id"], "motivo": "Sin texto extraíble; requiere OCR o revisión manual"})
            for unit in units:
                if not unit["texto"].strip():
                    gaps.append({"documento": doc["id"], "motivo": unit["ubicacion"] + ": sin texto; revisar imagen/OCR"})
                for fragment in re.split(r"\n+|(?<=[.;])\s+", unit["texto"]):
                    categories = [key for key, pattern in CATEGORIES.items() if re.search(pattern, _fold(fragment))]
                    if categories:
                        requirements.append({"id": f"R{len(requirements) + 1:04}", "categorias": categories,
                            "cita": fragment.strip(), "documento": doc["id"], "ubicacion": unit["ubicacion"],
                            "url": doc.get("url_final", doc["url"]), "sha256": doc["sha256"],
                            "estado": "CANDIDATO_POR_VALIDAR", "respuesta": PENDING})
        except (OSError, ValueError, ImportError, ET.ParseError, zipfile.BadZipFile) as exc:
            gaps.append({"documento": doc["id"], "motivo": str(exc)})
        except Exception as exc:
            gaps.append({"documento": doc["id"], "motivo": f"Error de extracción: {exc}"})
    found = {c for req in requirements for c in req["categorias"]}
    result = {"expediente": expediente, "analizado_en": now(), "requisitos": requirements,
              "faltantes": gaps, "categorias_no_identificadas": sorted(set(CATEGORIES) - found),
              "avisos": ["Extracción heurística, no dictamen de cumplimiento ni lectura exhaustiva.",
                         "Verificar tablas, imágenes, anexos y precedencia de adendas; no se resuelven contradicciones automáticamente.",
                         "El inventario público puede estar incompleto. No identificar un requisito no prueba que no exista."]}
    save(directory / "textos.json", texts)
    save(directory / "analisis.json", result)
    return result


def money(value):
    try:
        amount = Decimal(str(value))
    except (InvalidOperation, TypeError) as exc:
        raise ValueError("Importe/cantidad inválido; use números con punto decimal") from exc
    if not amount.is_finite() or amount < 0:
        raise ValueError("Importes y cantidades deben ser finitos y no negativos")
    return amount


def amount_text(value):
    return str(value.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP))


def economic(items):
    rows, total = [], Decimal(0)
    complete = bool(items)
    for item in items:
        row = {key: item.get(key, PENDING) for key in ("descripcion", "unidad", "cantidad", "precio_unitario", "impuesto_porcentaje", "fuente")}
        values = [item.get(k) for k in ("cantidad", "precio_unitario", "impuesto_porcentaje")]
        if any(v is None or v == "" for v in values):
            row.update(subtotal=PENDING, impuesto=PENDING, total=PENDING)
            complete = False
        else:
            qty, price, tax = map(money, values)
            subtotal = (qty * price).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
            tax_value = (subtotal * tax / 100).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
            row.update(subtotal=amount_text(subtotal), impuesto=amount_text(tax_value), total=amount_text(subtotal + tax_value))
            total += subtotal + tax_value
        rows.append(row)
    return rows, amount_text(total) if complete else PENDING


def cell(value):
    return str(value).replace("|", "\\|").replace("\n", " ").replace("\r", " ")


def csv_cell(value):
    value = str(value)
    return "'" + value if value.lstrip().startswith(("=", "+", "-", "@")) else value


def write_csv(path, rows, fields):
    with path.open("w", encoding="utf-8-sig", newline="") as output:
        writer = csv.DictWriter(output, fieldnames=fields, extrasaction="ignore")
        writer.writeheader()
        writer.writerows({key: csv_cell(row.get(key, "")) for key in fields} for row in rows)


def generate(expediente, oferta=None):
    """Genera Markdown y CSV editables; las respuestas son datos aportados, no cumplimiento certificado."""
    oferta = oferta or {}
    directory = folder(expediente)
    manifest = json.loads((directory / "manifest.json").read_text(encoding="utf-8"))
    analysis = analyze(expediente)
    responses = oferta.get("respuestas", {})
    known = {r["id"] for r in analysis["requisitos"]}
    if set(responses) - known:
        raise ValueError("La oferta contiene respuestas a requisitos inexistentes; revise analisis.json")
    rows, total = economic(oferta.get("items", []))
    version = "borrador-" + uuid.uuid4().hex[:12]
    destination = directory / version
    destination.mkdir()
    process = manifest["proceso"]
    lines = ["# BORRADOR DE PROPUESTA — REQUIERE REVISIÓN", "",
        f"Proceso/contrato: {cell(process.get('id_contrato') or process.get('id_proceso') or PENDING)}",
        f"Fuente: {cell(process['fuente'])} · Vigencia informada: {cell(process.get('vigencia', 'REVISAR'))}",
        f"Entidad: {cell(process.get('entidad') or PENDING)}", f"Objeto: {cell(process.get('objeto') or PENDING)}",
        f"Consulta del expediente: {manifest['creado_en']}",
        "", "Documento de trabajo: no constituye oferta presentada ni acredita habilitación. La vigencia debe reconfirmarse en el portal.",
        "", "## Proponente (información aportada)", cell(oferta.get("proponente") or PENDING),
        "", "## Propuesta técnica", cell(oferta.get("enfoque_tecnico") or PENDING),
        "", "### Respuesta a requisitos y criterios identificados",
        "Los extractos son candidatos por validar. Las respuestas aportadas requieren soportes y revisión de adendas.", "",
        "| ID / Categoría | Extracto del proceso | Fuente | Respuesta propuesta |",
        "| --- | --- | --- | --- |"]
    matrix = []
    for req in analysis["requisitos"]:
        response = responses.get(req["id"]) or PENDING
        reference = f"{req['documento']}, {req['ubicacion']}"
        lines.append("| " + " | ".join(map(cell, [req["id"] + " / " + ", ".join(req["categorias"]), req["cita"], reference, response])) + " |")
        matrix.append({**req, "categorias": ", ".join(req["categorias"]), "respuesta": response})
    if not matrix:
        lines.append("| — | No se identificaron requisitos verificables | — | PENDIENTE DE INFORMACIÓN |")
    lines += ["", "## Propuesta económica", f"Moneda: {cell(oferta.get('moneda') or PENDING)}",
              "Precios, cantidades e impuestos provienen de la oferta aportada; no se infieren del presupuesto oficial.", "",
              "| Descripción | Unidad | Cantidad | Precio unitario | Impuesto % | Total |",
              "| --- | --- | --- | --- | --- | --- |"]
    for row in rows:
        lines.append("| " + " | ".join(cell(row[k]) for k in ("descripcion", "unidad", "cantidad", "precio_unitario", "impuesto_porcentaje", "total")) + " |")
    if not rows:
        lines.append("| PENDIENTE DE INFORMACIÓN | — | — | — | — | — |")
    lines += [f"\nTotal calculado de los ítems aportados: {total}",
              "No certifica que se incluyan todos los ítems, tributos, costos o descuentos exigidos.",
              "", "## Propuesta comercial"]
    for key in ("plazo", "forma_pago", "validez_oferta", "garantias"):
        lines.append(f"- {key.replace('_', ' ')}: {cell(oferta.get(key) or PENDING)}")
    lines += ["", "## Pendientes y revisión", *[f"- {v}" for v in analysis["avisos"]],
              *[f"- {cell(g['documento'])}: {cell(g['motivo'])}" for g in analysis["faltantes"]],
              *[f"- {category}: no identificada en el texto extraído; verificar manualmente." for category in analysis["categorias_no_identificadas"]],
              "- Confirmar correspondencia de cada documento con el proceso, versión y adendas vigentes.",
              "- Completar soportes del proponente, anexos, firmas y todos los campos pendientes antes de presentar.",
              "", "## Documentos y trazabilidad",
              f"Descubrimiento: {cell(manifest['descubrimiento']['estado'])}; inventario completo no verificado.",
              cell(manifest['descubrimiento'].get('motivo') or manifest['descubrimiento'].get('aviso') or ''),
              f"Página consultada: {cell(process.get('url') or 'No aportada')}"]
    for doc in manifest["documentos"]:
        lines.append(f"- {doc['id']}: {doc['estado']} — {cell(doc['url'])} — SHA256: {doc.get('sha256', 'no disponible')} — {cell(doc.get('motivo', ''))}")
    (destination / "propuesta.md").write_text("\n".join(lines) + "\n", encoding="utf-8")
    save(destination / "oferta.json", oferta)
    save(destination / "analisis.json", analysis)
    write_csv(destination / "matriz_requisitos.csv", matrix,
              ["id", "categorias", "cita", "documento", "ubicacion", "url", "sha256", "estado", "respuesta"])
    write_csv(destination / "presupuesto.csv", rows,
              ["descripcion", "unidad", "cantidad", "precio_unitario", "impuesto_porcentaje", "subtotal", "impuesto", "total", "fuente"])
    return {"expediente": expediente, "directorio": str(destination),
            "archivos": [str(p) for p in destination.iterdir()], "total_items_aportados": total,
            "estado": "BORRADOR_REQUIERE_REVISION", "requisitos_candidatos": len(matrix)}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    sub = parser.add_subparsers(dest="command", required=True)
    p = sub.add_parser("obtener")
    p.add_argument("--proceso", required=True, help="JSON con el resultado seleccionado de búsqueda")
    p.add_argument("--url-documento", action="append", default=[])
    for name in ("analizar", "propuesta"):
        p = sub.add_parser(name)
        p.add_argument("--expediente", required=True)
        if name == "propuesta":
            p.add_argument("--oferta", help="JSON con información real aportada por el proponente")
    args = parser.parse_args()
    if args.command == "obtener":
        result = obtain(json.loads(Path(args.proceso).read_text(encoding="utf-8-sig")), args.url_documento)
    elif args.command == "analizar":
        result = analyze(args.expediente)
    else:
        result = generate(args.expediente, json.loads(Path(args.oferta).read_text(encoding="utf-8-sig")) if args.oferta else None)
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
